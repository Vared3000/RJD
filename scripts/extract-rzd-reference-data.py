"""Extract every source row/page and normalized candidates from the RZD reference folder.

The JSON output is intentionally written to server/tmp (gitignored).  It contains
personal data and must never be committed.  Run with Python 3.11+ and:
  pip install openpyxl xlrd pypdf
"""

from __future__ import annotations

import argparse
import hashlib
import json
import os
import re
import sys
from datetime import date, datetime
from pathlib import Path
from typing import Any, Iterable

extra_python_path = os.environ.get("RZD_IMPORT_PYTHON_PATH")
if extra_python_path:
    sys.path.insert(0, extra_python_path)
else:
    local_app_data = os.environ.get("LOCALAPPDATA")
    if local_app_data:
        sys.path.insert(0, str(Path(local_app_data) / "Temp" / "rzd-xlrd"))

import openpyxl
import xlrd
from pypdf import PdfReader


DPO_ALIASES = {
    "вос.сибир": "Восточно-Сибирская ДПО",
    "в-сибир": "Восточно-Сибирская ДПО",
    "восточно-сибир": "Восточно-Сибирская ДПО",
    "горьков": "Горьковская ДПО",
    "дальневост": "Дальневосточная ДПО",
    "д-вост": "Дальневосточная ДПО",
    "запсиб": "Западно-Сибирская ДПО",
    "западно-сибир": "Западно-Сибирская ДПО",
    "калининград": "Калининградская ДПО",
    "клнг": "Калининградская ДПО",
    "краснояр": "Красноярская ДПО",
    "куйбыш": "Куйбышевская ДПО",
    "москов": "Московская ДПО",
    "октябр": "Октябрьская ДПО",
    "приволж": "Приволжская ДПО",
    "свердлов": "Свердловская ДПО",
    "северо-кавказ": "Северо-Кавказская ДПО",
    "с-кав": "Северо-Кавказская ДПО",
    "северная": "Северная ДПО",
    "юго-вост": "Юго-Восточная ДПО",
    "ю-вост": "Юго-Восточная ДПО",
    "южно-урал": "Южно-Уральская ДПО",
    "ю-урал": "Южно-Уральская ДПО",
}

KNOWN_DPOS = sorted(set(DPO_ALIASES.values()))
DATE_RE = re.compile(r"(?<!\d)(\d{1,2})[./](\d{1,2})[./](20\d{2}|\d{2})(?!\d)")
PHONE_RE = re.compile(r"(?:\+7|8)[\s()\-]*\d{3}[\s()\-]*\d{3}[\s\-]*\d{2}[\s\-]*\d{2}")
PERSONNEL_RE = re.compile(r"^\d{5,12}$")
SIZE_VALUE_RE = re.compile(r"\d+(?:[.,]\d+)?")
MONTHS = {
    "январ": 1,
    "феврал": 2,
    "март": 3,
    "апрел": 4,
    "ма(?:й|я)": 5,
    "июн": 6,
    "июл": 7,
    "август": 8,
    "сентябр": 9,
    "октябр": 10,
    "ноябр": 11,
    "декабр": 12,
}


def clean_text(value: Any) -> str:
    return re.sub(r"\s+", " ", str(value or "")).strip()


def lower(value: Any) -> str:
    return clean_text(value).lower().replace("ё", "е")


def json_value(value: Any) -> Any:
    if value is None:
        return None
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    if isinstance(value, float):
        return int(value) if value.is_integer() else value
    if isinstance(value, (str, int, bool)):
        return value
    return str(value)


def trim(values: list[Any]) -> list[Any]:
    values = list(values)
    while values and values[-1] in (None, ""):
        values.pop()
    return values


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def source_key(file_hash: str, locator: str) -> str:
    return hashlib.sha256(f"{file_hash}:{locator}".encode("utf-8")).hexdigest()


def parse_date(text: str) -> str | None:
    dates = []
    for day, month, year in DATE_RE.findall(text):
        year = f"20{year}" if len(year) == 2 else year
        try:
            dates.append(date(int(year), int(month), int(day)))
        except ValueError:
            pass
    lowered = lower(text)
    for prefix, month in MONTHS.items():
        match = re.search(rf"{prefix}[а-я]*\s+(20\d{{2}})", lowered)
        if match:
            # An act labelled by month represents the closing reporting date.
            next_month = date(int(match.group(1)) + (month == 12), month % 12 + 1, 1)
            dates.append(date.fromordinal(next_month.toordinal() - 1))
    return max(dates).isoformat() if dates else None


def iso_date(day: str, month: str, year: str) -> str | None:
    year = f"20{year}" if len(year) == 2 else year
    try:
        return date(int(year), int(month), int(day)).isoformat()
    except ValueError:
        return None


def detect_dpo(*texts: str) -> str | None:
    haystack = " ".join(lower(text) for text in texts)
    for alias, canonical in DPO_ALIASES.items():
        if alias in haystack:
            return canonical
    return None


def normalize_name(value: Any) -> str | None:
    text = clean_text(value)
    text = re.sub(r"\([^)]*(?:комплект|замен|увол|размер)[^)]*\)", "", text, flags=re.I)
    text = clean_text(text)
    if not text or any(token in lower(text) for token in ("фио", "ф.и.о", "работника заказчика")):
        return None
    words = re.findall(r"[А-ЯЁA-Z][а-яёa-z-]+", text)
    if 2 <= len(words) <= 5 and len(" ".join(words)) >= 8:
        return " ".join(words)
    return None


def last_person_name(value: Any) -> str | None:
    words = re.findall(r"[А-ЯЁA-Z][а-яёa-z-]+", clean_text(value))
    return " ".join(words[-3:]) if len(words) >= 3 else None


def numeric(value: Any) -> float | None:
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        return float(value)
    text = clean_text(value).replace(" ", "").replace(",", ".")
    if re.fullmatch(r"-?\d+(?:\.\d+)?", text):
        return float(text)
    return None


def split_measurements(value: Any) -> list[str]:
    text = clean_text(value)
    if not text:
        return []
    return list(dict.fromkeys(match.replace(",", ".") for match in SIZE_VALUE_RE.findall(text)))


def classify_size_type(item_name: str) -> tuple[str | None, bool]:
    text = lower(item_name)
    if any(word in text for word in ("обув", "туфл", "ботин", "сапог", "полубот")):
        return "shoe", False
    if any(word in text for word in ("головн", "шапк", "кепк", "пилотк", "фуражк")):
        return "headwear", False
    if "ремен" in text:
        return "belt", False
    if any(word in text for word in ("перчат", "вареж")):
        return "gloves", False
    apparel = (
        "жакет", "пиджак", "юбк", "брюк", "блуз", "сороч", "рубаш", "жилет",
        "пальто", "плащ", "куртк", "костюм", "свитер", "джемпер", "плать",
    )
    if any(word in text for word in apparel):
        return "clothing", True
    return None, False


def column(headers: list[str], *tokens: str, exclude: tuple[str, ...] = ()) -> int | None:
    for index, header in enumerate(headers):
        if all(token in header for token in tokens) and not any(token in header for token in exclude):
            return index
    return None


def cell(values: list[Any], index: int | None) -> Any:
    return values[index] if index is not None and index < len(values) else None


def read_xlsx(path: Path) -> Iterable[tuple[str, list[dict[str, Any]]]]:
    values_book = openpyxl.load_workbook(path, read_only=True, data_only=True)
    formula_book = openpyxl.load_workbook(path, read_only=True, data_only=False)
    try:
        for values_sheet, formula_sheet in zip(values_book.worksheets, formula_book.worksheets):
            rows = []
            for row_number, (value_row, formula_row) in enumerate(
                zip(values_sheet.iter_rows(), formula_sheet.iter_rows()), start=1
            ):
                values = trim([json_value(item.value) for item in value_row])
                formulas = {
                    str(index): item.value
                    for index, item in enumerate(formula_row, start=1)
                    if isinstance(item.value, str) and item.value.startswith("=")
                }
                if values or formulas:
                    rows.append({"rowNumber": row_number, "values": values, "formulas": formulas})
            yield values_sheet.title, rows
    finally:
        values_book.close()
        formula_book.close()


def read_xls(path: Path) -> Iterable[tuple[str, list[dict[str, Any]]]]:
    book = xlrd.open_workbook(path, on_demand=True)
    try:
        for sheet in book.sheets():
            rows = []
            for row_number in range(sheet.nrows):
                values = []
                for col_number in range(sheet.ncols):
                    item = sheet.cell(row_number, col_number)
                    value: Any = item.value
                    if item.ctype == xlrd.XL_CELL_DATE:
                        value = xlrd.xldate_as_datetime(value, book.datemode).isoformat()
                    values.append(json_value(value))
                values = trim(values)
                if values:
                    rows.append({"rowNumber": row_number + 1, "values": values, "formulas": {}})
            yield sheet.name, rows
    finally:
        book.release_resources()


def extract_dpo_details(
    relative_file: str,
    file_hash: str,
    sheet_name: str,
    rows: list[dict[str, Any]],
    dpo: str | None,
) -> dict[str, Any] | None:
    if not dpo:
        return None

    first_rows = rows[:40]
    row_texts = [
        clean_text(" ".join(clean_text(value) for value in row["values"]))
        for row in first_rows
    ]
    header_text = " ".join(row_texts)
    full_name = None
    address = None
    for line in row_texts[:15]:
        if (
            detect_dpo(line) == dpo
            and "дирекция пассажирских обустройств" in lower(line)
            and re.search(r"\b\d{6}\b", line)
        ):
            postal = re.search(r"\b\d{6}\b", line)
            full_name = clean_text(line[: postal.start()].rstrip(" ,"))
            address = clean_text(line[postal.start() :])
            break

    business_unit_code = None
    okpo = None
    for row in first_rows[:15]:
        values = row["values"]
        for index, value in enumerate(values):
            label = lower(value)
            following = clean_text(values[index + 1]) if index + 1 < len(values) else ""
            if label == "бе" and following and following != "-":
                business_unit_code = following
            if "по окпо" in label and following:
                okpo = following

    additional_agreement_number = None
    additional_agreement_date = None
    agreement = re.search(
        rf"дополнительн\w*\s+соглашени\w*\s+(?:№\s*)?([0-9А-ЯA-ZЁа-яё/.\-]+)"
        rf"\s+от\s+{DATE_RE.pattern}",
        header_text,
        flags=re.I,
    )
    if agreement:
        additional_agreement_number = agreement.group(1)
        additional_agreement_date = iso_date(
            agreement.group(2), agreement.group(3), agreement.group(4)
        )

    contract_number = None
    contract_date = None
    contract_after_number = re.search(
        rf"(?:к\s+)?договор\w*\s+№\s*([0-9А-ЯA-ZЁа-яё/.\-]+)"
        rf"\s+от\s+{DATE_RE.pattern}",
        header_text,
        flags=re.I,
    )
    contract_after_date = re.search(
        rf"договор\w*\s+от\s+{DATE_RE.pattern}(?:\s+года?)?\s+№\s*"
        rf"([0-9А-ЯA-ZЁа-яё/.\-]+)",
        header_text,
        flags=re.I,
    )
    if contract_after_number:
        contract_number = contract_after_number.group(1)
        contract_date = iso_date(
            contract_after_number.group(2),
            contract_after_number.group(3),
            contract_after_number.group(4),
        )
    elif contract_after_date:
        contract_date = iso_date(
            contract_after_date.group(1),
            contract_after_date.group(2),
            contract_after_date.group(3),
        )
        contract_number = contract_after_date.group(4)

    director_full_name = None
    director_basis = None
    representative_index = next(
        (
            index
            for index, line in enumerate(row_texts)
            if "представитель заказчика" in lower(line)
        ),
        None,
    )
    if representative_index is not None:
        for line in row_texts[representative_index : representative_index + 5]:
            candidate = last_person_name(line)
            if candidate and "дирекц" not in lower(candidate):
                director_full_name = candidate
                break
        basis_match = re.search(
            r"действующ\w*\s+на\s+основании\s*(.+)",
            " ".join(row_texts[representative_index : representative_index + 6]),
            flags=re.I,
        )
        if basis_match:
            director_basis = clean_text(
                re.split(r"\(\s*вид документа|составили настоящий акт", basis_match.group(1), flags=re.I)[0]
            )
    if not director_full_name:
        customer_side = re.search(
            r"Заказчик.+?в лице (.+?),\s*действующ\w*\s+на\s+основании\s+"
            r"(.+?)(?:,\s*с одной стороны|,\s*и\s+Общество)",
            header_text,
            flags=re.I,
        )
        if customer_side:
            director_full_name = last_person_name(customer_side.group(1))
            director_basis = clean_text(customer_side.group(2))

    values = {
        "type": "dpo",
        "sourceKey": source_key(file_hash, f"sheet:{sheet_name}:dpo-details"),
        "sourceFile": relative_file,
        "sheetName": sheet_name,
        "dpo": dpo,
        "fullName": full_name,
        "address": address,
        "okpo": okpo,
        "businessUnitCode": business_unit_code,
        "directorFullName": director_full_name,
        "directorBasis": director_basis,
        "contractNumber": contract_number,
        "contractDate": contract_date,
        "additionalAgreementNumber": additional_agreement_number,
        "additionalAgreementDate": additional_agreement_date,
    }
    return values if any(value for key, value in values.items() if key not in {
        "type", "sourceKey", "sourceFile", "sheetName", "dpo"
    }) else None


def analyze_sheet(
    relative_file: str,
    file_hash: str,
    sheet_name: str,
    rows: list[dict[str, Any]],
    file_dpo: str | None,
) -> list[dict[str, Any]]:
    candidates: list[dict[str, Any]] = []
    row_values = [row["values"] for row in rows]
    sheet_text = " ".join(clean_text(value) for values in row_values[:15] for value in values)
    # A few workbooks are copies stored in another road's folder; the sheet
    # heading is therefore more authoritative than the path.
    dpo = detect_dpo(sheet_name, sheet_text) or file_dpo
    document_date = parse_date(f"{relative_file} {sheet_text}")
    dpo_details = extract_dpo_details(relative_file, file_hash, sheet_name, rows, dpo)
    if dpo_details:
        candidates.append(dpo_details)

    # One sheet may contain one or more repeated tables. Each recognized header
    # governs rows until the next recognized header.
    header_positions = []
    for index, values in enumerate(row_values):
        headers = [lower(value) for value in values]
        joined = " ".join(headers)
        if (
            ("фио" in joined or "ф.и.о" in joined or "должность , фио" in joined)
            and ("табель" in joined)
        ) or (
            "наименование форменной одежды" in joined
            and ("должность" in joined or "фио" in joined)
        ) or (
            "наименование видов" in joined
            and "выполненных работ" in joined
        ):
            header_positions.append(index)

    for header_order, header_index in enumerate(header_positions):
        end_index = header_positions[header_order + 1] if header_order + 1 < len(header_positions) else len(rows)
        headers = [lower(value) for value in row_values[header_index]]
        data_start_index = header_index + 1
        if (
            "наименование видов" in " ".join(headers)
            and "выполненных работ" in " ".join(headers)
            and data_start_index < len(row_values)
        ):
            secondary_headers = [lower(value) for value in row_values[data_start_index]]
            width = max(len(headers), len(secondary_headers))
            headers = [
                " ".join(
                    value
                    for value in (
                        headers[index] if index < len(headers) else "",
                        secondary_headers[index] if index < len(secondary_headers) else "",
                    )
                    if value
                )
                for index in range(width)
            ]
            data_start_index += 1
        fio_idx = column(headers, "фио")
        if fio_idx is None:
            fio_idx = column(headers, "должность", "фио")
        personnel_idx = column(headers, "табель")
        position_idx = column(headers, "должность", exclude=("фио",))
        combined_position_idx = column(headers, "должность", "фио")
        item_idx = column(headers, "наименование", "форменной")
        fpu_item_idx = column(headers, "наименование видов", "выполненных работ")
        if item_idx is None:
            item_idx = fpu_item_idx
        form_type = "fpu-26" if fpu_item_idx is not None else None
        quantity_idx = column(headers, "кол-во")
        if quantity_idx is None:
            quantity_idx = column(headers, "количество")
        coverage_days_idx = column(headers, "дней", "обеспеч")
        clothing_idx = column(headers, "размер", "одежд")
        height_idx = column(headers, "рост")
        shoe_idx = column(headers, "размер", "обув")
        headwear_idx = column(headers, "размер", "голов")
        belt_idx = column(headers, "ремень")
        gloves_idx = column(headers, "перчат")
        unit_idx = column(headers, "ед")
        no_vat_idx = column(headers, "без ндс", exclude=("итого", "сумма"))
        with_vat_idx = column(headers, "с ндс", exclude=("итого", "сумма"))
        vat_amount_idx = column(headers, "сумма", "ндс")
        if vat_amount_idx is None:
            vat_amount_idx = next(
                (
                    index
                    for index, header in enumerate(headers)
                    if clean_text(header).startswith("ндс")
                ),
                None,
            )
        displayed_price_no_vat_idx = column(headers, "цена за единицу")
        total_without_vat_idx = column(headers, "итого стоимость", "без ндс")
        row_total_with_vat_idx = column(headers, "сумма с ндс")
        if row_total_with_vat_idx is None:
            row_total_with_vat_idx = column(headers, "итого стоимость", "с ндс")
        # В приложении 1.7 заголовок «Цена ... без НДС» объединяет две
        # колонки: слева цена за единицу, справа сумма строки. В выгрузке
        # merged-cell заголовок остаётся только в правой колонке.
        unit_price_no_vat_idx = no_vat_idx
        subtotal_no_vat_idx = no_vat_idx
        if (
            displayed_price_no_vat_idx is not None
            and displayed_price_no_vat_idx > 0
            and not clean_text(headers[displayed_price_no_vat_idx - 1])
        ):
            unit_price_no_vat_idx = displayed_price_no_vat_idx - 1
        if (
            no_vat_idx is not None
            and no_vat_idx > 0
            and not clean_text(headers[no_vat_idx - 1])
        ):
            unit_price_no_vat_idx = no_vat_idx - 1
            subtotal_no_vat_idx = no_vat_idx

        current_employee: dict[str, Any] | None = None
        current_position: str | None = None
        for local_index in range(data_start_index, end_index):
            row = rows[local_index]
            values = row["values"]
            joined = lower(" ".join(clean_text(value) for value in values))
            if (
                "всего:" in joined
                or "итого:" in joined
                or joined.strip().startswith("итого")
            ):
                break
            if any(token in joined for token in ("исполнитель:", "заказчик:")):
                current_employee = None
                continue

            personnel_raw = clean_text(cell(values, personnel_idx))
            personnel = personnel_raw if PERSONNEL_RE.fullmatch(personnel_raw) else None
            raw_name = cell(values, fio_idx if fio_idx is not None else combined_position_idx)
            name = normalize_name(raw_name)

            if combined_position_idx is not None and not personnel and raw_name:
                possible_position = clean_text(raw_name)
                if not normalize_name(possible_position) and len(possible_position) > 3:
                    current_position = possible_position
            if position_idx is not None and cell(values, position_idx):
                current_position = clean_text(cell(values, position_idx))

            if name and (personnel or fio_idx is not None):
                current_employee = {
                    "fullName": name,
                    "personnelNumber": personnel,
                    "position": current_position,
                    "dpo": dpo,
                    "phone": None,
                    "hireDate": None,
                    "terminationDate": None,
                    "measurements": {},
                }
                for size_type, index in (
                    ("clothing", clothing_idx),
                    ("height", height_idx),
                    ("shoe", shoe_idx),
                    ("headwear", headwear_idx),
                    ("belt", belt_idx),
                    ("gloves", gloves_idx),
                ):
                    measured = split_measurements(cell(values, index))
                    if measured:
                        current_employee["measurements"][size_type] = measured
                phones = PHONE_RE.findall(" ".join(clean_text(value) for value in values))
                if phones:
                    current_employee["phone"] = clean_text(phones[0])
                candidates.append(
                    {
                        "type": "employee",
                        "sourceKey": source_key(
                            file_hash, f"sheet:{sheet_name}:row:{row['rowNumber']}"
                        ),
                        "rowNumber": row["rowNumber"],
                        **current_employee,
                    }
                )

            item_name = clean_text(cell(values, item_idx))
            if item_name and "наименование" not in lower(item_name) and len(item_name) > 2:
                size_type, requires_height = classify_size_type(item_name)
                source_formulas = {}
                for field, index in (
                    ("displayedPriceWithoutVat", displayed_price_no_vat_idx),
                    ("costWithoutVat", subtotal_no_vat_idx),
                    ("totalWithoutVat", total_without_vat_idx),
                    ("vatAmount", vat_amount_idx),
                    (
                        "totalWithVat",
                        row_total_with_vat_idx
                        if row_total_with_vat_idx is not None
                        else with_vat_idx,
                    ),
                ):
                    formula_value = row["formulas"].get(str(index + 1)) if index is not None else None
                    if formula_value and len(formula_value) > 1:
                        source_formulas[field] = formula_value[1:]
                candidates.append(
                    {
                        "type": "nomenclature",
                        "sourceKey": source_key(
                            file_hash, f"sheet:{sheet_name}:row:{row['rowNumber']}"
                        ),
                        "rowNumber": row["rowNumber"],
                        "name": item_name,
                        "unit": clean_text(cell(values, unit_idx)) or "шт.",
                        "sizeType": size_type,
                        "requiresHeightSize": requires_height,
                        "position": current_position,
                        "quantity": numeric(cell(values, quantity_idx)),
                        "coverageDays": numeric(cell(values, coverage_days_idx)),
                        "employee": current_employee,
                        "dpo": dpo,
                        "effectiveDate": document_date,
                        "formType": form_type,
                        "priceWithoutVat": numeric(cell(values, unit_price_no_vat_idx)),
                        "displayedPriceWithoutVat": numeric(
                            cell(values, displayed_price_no_vat_idx)
                        ),
                        "subtotalWithoutVat": numeric(cell(values, subtotal_no_vat_idx)),
                        "totalWithoutVat": numeric(cell(values, total_without_vat_idx)),
                        "vatAmount": numeric(cell(values, vat_amount_idx)),
                        "priceWithVat": (
                            None
                            if form_type == "fpu-26"
                            else numeric(cell(values, with_vat_idx))
                        ),
                        "totalWithVat": numeric(
                            cell(
                                values,
                                row_total_with_vat_idx
                                if row_total_with_vat_idx is not None
                                else with_vat_idx,
                            )
                        ),
                        "sourceFormulas": source_formulas,
                    }
                )

    # Personal-card prose often places all data in merged cells rather than a table.
    card_text = clean_text(" ".join(sheet_text.split()))
    card_match = re.search(
        r"Работник\s+Заказчика:\s*(.+?)\s+таб\.?\s*№\s*([0-9]+)(?:\s+должность:\s*(.+?))?(?=\s+(?:Индивидуальные|Размер|Дата|$))",
        card_text,
        flags=re.I,
    )
    if card_match:
        name = normalize_name(card_match.group(1))
        if name:
            measurements: dict[str, list[str]] = {}
            pair = re.search(r"Индивидуальные размеры одежды:\s*([0-9, ]+)\s*/\s*([0-9, ]+)", card_text, re.I)
            if pair:
                measurements["clothing"] = split_measurements(pair.group(1))
                measurements["height"] = split_measurements(pair.group(2))
            for size_type, pattern in (
                ("shoe", r"размер обуви[:\s]+([0-9,./ ]+)"),
                ("headwear", r"размер головного убора[:\s]+([0-9,./ ]+)"),
                ("belt", r"размер ремня[:\s]+([0-9,./ ]+)"),
                ("gloves", r"размер перчат(?:ок|ки)[:\s]+([0-9,./ ]+)"),
            ):
                match = re.search(pattern, card_text, re.I)
                if match:
                    measurements[size_type] = split_measurements(match.group(1))
            candidates.append(
                {
                    "type": "employee",
                    "sourceKey": source_key(file_hash, f"sheet:{sheet_name}:personal-card"),
                    "fullName": name,
                    "personnelNumber": card_match.group(2),
                    "position": clean_text(card_match.group(3)) or None,
                    "dpo": dpo,
                    "phone": (PHONE_RE.search(card_text).group(0) if PHONE_RE.search(card_text) else None),
                    "hireDate": None,
                    "terminationDate": None,
                    "measurements": measurements,
                }
            )
    for candidate in candidates:
        candidate["sourceFile"] = relative_file
        candidate["sheetName"] = sheet_name
    return candidates


def extract(source_root: Path) -> dict[str, Any]:
    files: list[dict[str, Any]] = []
    candidates: list[dict[str, Any]] = []
    errors: list[dict[str, str]] = []

    for path in sorted(item for item in source_root.rglob("*") if item.is_file() and not item.name.startswith("~$")):
        relative = path.relative_to(source_root).as_posix()
        file_hash = sha256_file(path)
        file_record: dict[str, Any] = {
            "path": relative,
            "hash": file_hash,
            "size": path.stat().st_size,
            "extension": path.suffix.lower(),
        }
        file_dpo = detect_dpo(relative)
        try:
            if path.suffix.lower() in (".xlsx", ".xls"):
                reader = read_xlsx if path.suffix.lower() == ".xlsx" else read_xls
                sheets = []
                for sheet_name, rows in reader(path):
                    sheets.append({"name": sheet_name, "rows": rows})
                    candidates.extend(analyze_sheet(relative, file_hash, sheet_name, rows, file_dpo))
                file_record["sheets"] = sheets
            elif path.suffix.lower() == ".pdf":
                pdf = PdfReader(path)
                file_record["pages"] = [
                    {"pageNumber": index, "text": page.extract_text() or ""}
                    for index, page in enumerate(pdf.pages, start=1)
                ]
            else:
                file_record["metadataOnly"] = True
        except Exception as error:  # preserve file metadata even when a parser fails
            file_record["parseError"] = repr(error)
            errors.append({"path": relative, "error": repr(error)})
        files.append(file_record)

    return {
        "formatVersion": 1,
        "generatedAt": datetime.now().astimezone().isoformat(),
        "sourceRoot": str(source_root.resolve()),
        "knownDpos": KNOWN_DPOS,
        "files": files,
        "candidates": candidates,
        "errors": errors,
    }


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("source", type=Path)
    parser.add_argument("output", type=Path)
    args = parser.parse_args()
    result = extract(args.source)
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(json.dumps(result, ensure_ascii=False), encoding="utf-8")
    row_count = sum(
        len(sheet["rows"])
        for file in result["files"]
        for sheet in file.get("sheets", [])
    )
    page_count = sum(len(file.get("pages", [])) for file in result["files"])
    print(
        f"files={len(result['files'])} rows={row_count} pages={page_count} "
        f"candidates={len(result['candidates'])} errors={len(result['errors'])} "
        f"output={args.output}"
    )
    if result["errors"]:
        raise SystemExit(2)


if __name__ == "__main__":
    main()
